"""Monte Carlo over the IP valuation chain (same mechanics as the Scenarios sheet).

The deterministic model gives one number per story. This gives the distribution of
value per share when the four decisive assumptions are uncertain, and answers the
question a bet actually needs answered: P(value > price), the tails, and how big a
position the edge justifies (Kelly).

Everything here mirrors IP_valuation_model.xlsx — same inputs, same bridge, same
terminal mechanics. Run: python3 montecarlo.py  (writes mc.json for the post's d3 chart)
"""
import json
import numpy as np

rng = np.random.default_rng(2026)
N = 20000

# ---------------- fixed inputs, read from the workbook's Inputs sheet so the two can't drift ----------------
from openpyxl import load_workbook
I = {r[0].value: r[1].value for r in load_workbook("IP_valuation_model.xlsx", data_only=True)["Inputs"].iter_rows(min_row=5)
     if r[0].value and isinstance(r[1].value, (int, float))}
price, shares = float(I["Share price ($)"]), float(I["Diluted shares (millions)"])
net_debt = I["Gross debt"] - I["Cash and temporary investments"] + next(v for k, v in I.items() if "debt-like" in k.lower())
rev_2025_na, rev_2025_emea = float(I["2025 pro forma net sales — Packaging Solutions NA"]), float(I["2025 pro forma net sales — Packaging Solutions EMEA"])
rev_2026_h1 = float(I["H1 2026 net sales"])
na_26, emea_26 = float(I["2026 adj. EBITDA guidance — NA (midpoint)"]), float(I["2026 adj. EBITDA guidance — EMEA (midpoint)"])
corp = (I["2026 adj. EBITDA guidance — low"] + I["2026 adj. EBITDA guidance — high"]) / 2 - na_26 - emea_26
price_index_ytd = float(I["2026 containerboard index increase recognized YTD ($/ton, linerboard)"])
price_sens, price_in_2026 = float(I["EBITDA per $1/ton of price (annualized, NA)"]), float(I["Price benefit already inside 2026 NA guidance"])
contrib_margin = I["Contribution margin on incremental NA volume"]
costout_27, infl_27 = float(I["Incremental 80/20 cost-out 2027 (NA + EMEA)"]), float(I["Cost inflation 2027 (OCC, energy, labor, freight)"])
emea_27_delta = float(I["EMEA EBITDA change 2027 (box-price catch-up + cost-out, net)"])
rev_growth_lt, da_pct = I["Revenue growth 2028–2030 (nominal)"], I["D&A as % of revenue 2027+"]
nwc_pct, tax = I["Net working capital investment, % of incremental revenue"], I["Cash tax rate"]
capex_2027, sep_costs = float(I["Capex 2027"]), float(I["Separation / transaction costs expected 2026–27"])
g = I["Terminal growth rate"]

rev_na_26 = rev_2026_h1 * 2 * rev_2025_na / (rev_2025_na + rev_2025_emea)
rev_emea_26 = rev_2026_h1 * 2 * rev_2025_emea / (rev_2025_na + rev_2025_emea)
rev_26 = rev_na_26 + rev_emea_26

# ---------------- the uncertain assumptions, with a cycle ----------------
# Independent draws understate the tails: in a downturn price does NOT stick, demand falls,
# margins compress and the multiple shrinks together. So first draw the regime, then draw the
# assumptions conditional on it.  P(down-cycle over the horizon) is itself a judgment: 30%.
p_down = 0.30
down = rng.random(N) < p_down
# 1. Share of the $170/ton that sticks through 2027. Jason: "more like half, lag is months."
stick = np.where(down, rng.beta(2, 5, N), rng.beta(5, 4, N))          # down: mode ~0.2; base: mode ~0.57
# 2. Mid-cycle NA EBITDA margin: SW 13% floor ... PKG 21% ceiling.
na_margin = np.where(down, rng.triangular(0.11, 0.135, 0.16, N), rng.triangular(0.14, 0.175, 0.21, N))
# 3. Exit EV/EBITDA multiple.
mult = np.where(down, rng.triangular(6.0, 7.0, 8.5, N), rng.triangular(7.0, 8.5, 11.0, N))
# 4. WACC: normal around 7.4%, 50 bp sd, floored well above g.
wacc = np.clip(rng.normal(0.074, 0.005, N), 0.055, 0.10)
# 2027 volume: base +1% +/- 1.5%; down-cycle -2.5% +/- 1.5% (TTOBMA's -2% forecast lives here).
vol = np.where(down, rng.normal(-0.025, 0.015, N), rng.normal(0.01, 0.015, N))
# EMEA mid-cycle margin and the 2027 catch-up also depend on the regime.
emea_margin_mid = np.where(down, rng.triangular(0.08, 0.095, 0.11, N), rng.triangular(0.10, 0.115, 0.13, N))
emea_27_delta = np.where(down, rng.normal(0.0, 75.0, N), rng.normal(150.0, 75.0, N))
# Capex intensity 2028+: the 'assets in the ground' lever.
capex_pct = rng.triangular(0.060, 0.065, 0.078, N)
# Model risk: the structure itself is wrong by some amount we cannot see. Lognormal haircut/uplift on
# the terminal value with 15% sd — Damodaran's margin of safety, made explicit instead of applied at the end.
MODEL_NOISE = 0.15
model_noise = np.exp(rng.normal(0.0, MODEL_NOISE, N))

# ---------------- the chain (vectorized, same as Scenarios sheet) ----------------
price_delta = stick * price_index_ytd * price_sens - price_in_2026
rev_na_27 = rev_na_26 * (1 + vol) + price_delta
rev_emea_27 = rev_emea_26 * (1 + rev_growth_lt)
na_27 = na_26 + price_delta + rev_na_26 * vol * contrib_margin + (costout_27 - infl_27) * rev_na_26 / rev_26
emea_27 = emea_26 + emea_27_delta + (costout_27 - infl_27) * rev_emea_26 / rev_26
ebitda_27 = na_27 + emea_27 + corp
rev_27 = rev_na_27 + rev_emea_27
grow3 = (1 + rev_growth_lt) ** 3
rev_30 = rev_27 * grow3
ebitda_30 = rev_na_27 * grow3 * na_margin + rev_emea_27 * grow3 * emea_margin_mid + corp
ufcf_30 = (ebitda_30 - rev_30 * da_pct) * (1 - tax) + rev_30 * da_pct - rev_30 * capex_pct - rev_30 * rev_growth_lt / (1 + rev_growth_lt) * nwc_pct
ufcf_27 = (ebitda_27 - rev_27 * da_pct) * (1 - tax) + rev_27 * da_pct - capex_2027 - sep_costs * 0.5
pv_interim = (ufcf_27 / (1 + wacc) ** 0.7
              + (ufcf_27 * 2 / 3 + ufcf_30 / 3) / (1 + wacc) ** 1.7
              + (ufcf_27 / 3 + ufcf_30 * 2 / 3) / (1 + wacc) ** 2.7
              + ufcf_30 / (1 + wacc) ** 3.7)
tv = 0.5 * (ufcf_30 * (1 + g) / (wacc - g) + ebitda_30 * mult) * model_noise
ev = pv_interim + tv / (1 + wacc) ** 3.7
value = (ev - net_debt) / shares

# ---------------- what a bet needs to know ----------------
p_above = float((value > price).mean())
pct = {k: float(np.percentile(value, k)) for k in (5, 10, 25, 50, 75, 90, 95)}
ev_value = float(value.mean())
# Treat it as a bet that resolves over ~3 years: return r = value/price - 1 (ignoring the 5.8% dividend,
# which is upside if it survives and a cut if it doesn't). Growth-optimal (Kelly) fraction solves
# max_f E[log(1 + f r)] numerically; we report half of it, which is what practitioners actually size to.
r = value / price - 1
fs = np.linspace(0, 1, 201)
growth = np.array([np.mean(np.log1p(np.clip(f * r, -0.999, None))) for f in fs])
kelly = float(fs[growth.argmax()])
half_kelly = kelly / 2

# what drives it: rank correlation of each input with value
def spearman(a, b):
    ra, rb = a.argsort().argsort(), b.argsort().argsort()
    return float(np.corrcoef(ra, rb)[0, 1])
drivers = {"down-cycle regime": spearman(down.astype(float), value), "price sticks": spearman(stick, value), "NA mid-cycle margin": spearman(na_margin, value),
           "exit multiple": spearman(mult, value), "WACC": spearman(wacc, value),
           "2027 volume": spearman(vol, value), "capex intensity": spearman(capex_pct, value), "model noise": spearman(model_noise, value)}

hist, edges = np.histogram(value, bins=np.arange(0, 100, 2.5))
out = {"n": N, "price": price, "p_down": p_down, "p_above_price_given_down": float((value[down] > price).mean()), "p_above_price_given_base": float((value[~down] > price).mean()),
       "mean_down": float(value[down].mean()), "mean_base": float(value[~down].mean()), "p_above_price": p_above, "mean": ev_value, "percentiles": pct,
       "kelly": kelly, "half_kelly": half_kelly, "drivers": drivers,
       "hist": {"edges": edges.tolist(), "counts": hist.tolist()},
       "model_noise_sd": MODEL_NOISE, "p_loss_half": float((value < price * 0.5).mean()), "p_double": float((value > price * 2).mean())}
json.dump(out, open("mc.json", "w"), indent=1)
print(f"P(value > ${price}) = {p_above:.1%}   mean ${ev_value:.1f}   median ${pct[50]:.1f}")
print("percentiles:", {k: round(v, 1) for k, v in pct.items()})
print(f"P(value < half the price) = {out['p_loss_half']:.1%}   P(value > 2x) = {out['p_double']:.1%}")
print(f"Kelly {kelly:.2f}  half-Kelly {half_kelly:.2f}")
print(f"given down-cycle ({p_down:.0%}): P(>price) {out['p_above_price_given_down']:.1%}, mean ${out['mean_down']:.1f};  given base: P(>price) {out['p_above_price_given_base']:.1%}, mean ${out['mean_base']:.1f}")
print("drivers (rank corr):", {k: round(v, 2) for k, v in drivers.items()})
