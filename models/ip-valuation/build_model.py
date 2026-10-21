"""International Paper (NYSE: IP) valuation model — WSP mechanics, Damodaran framing.
Every assumption lives on Inputs with a provenance tag; all other sheets are formulas.
Colors: blue = hardcoded input, black = formula, green = cross-sheet link, yellow fill = floor judgment to overwrite.
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment

wb = Workbook()
FONT = "Arial"
BLUE = Font(name=FONT, color="0000FF", size=10)
BLACK = Font(name=FONT, color="000000", size=10)
GREEN = Font(name=FONT, color="008000", size=10)
BOLD = Font(name=FONT, bold=True, size=10)
H1 = Font(name=FONT, bold=True, size=13)
H2 = Font(name=FONT, bold=True, size=11, color="FFFFFF")
GREY = Font(name=FONT, italic=True, size=9, color="666666")
YELLOW = PatternFill("solid", fgColor="FFFF00")
HEADFILL = PatternFill("solid", fgColor="1F3A5F")
SUBFILL = PatternFill("solid", fgColor="DCE6F1")
THIN = Side(style="thin", color="999999")
TOP = Border(top=THIN)
USD = '$#,##0;($#,##0);-'
USD1 = '$#,##0.0;($#,##0.0);-'
USD2 = '$#,##0.00;($#,##0.00);-'
PCT = '0.0%;(0.0%);-'
MULT = '0.0"x"'
NUM = '#,##0;(#,##0);-'
NUM1 = '#,##0.0;(#,##0.0);-'

def setw(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

def header(ws, row, text, span=8):
    ws.cell(row=row, column=1, value=text).font = H2
    for c in range(1, span + 1):
        ws.cell(row=row, column=c).fill = HEADFILL

def sub(ws, row, text, span=8):
    ws.cell(row=row, column=1, value=text).font = BOLD
    for c in range(1, span + 1):
        ws.cell(row=row, column=c).fill = SUBFILL

# =====================================================================================
# INPUTS
# =====================================================================================
ws = wb.active
ws.title = "Inputs"
setw(ws, [46, 14, 12, 16, 70])
ws["A1"] = "International Paper (NYSE: IP) — valuation inputs"; ws["A1"].font = H1
ws["A2"] = "All figures $ millions unless noted. Blue = hardcoded input. Yellow = floor judgment: overwrite with your own view. Black = formula."; ws["A2"].font = GREY
ws["A3"] = "Provenance tags: REPORTED (company filing/release) · GUIDANCE (company outlook) · MARKET (quoted price/index) · JUDGMENT (this model's call — check it)."; ws["A3"].font = GREY

LAB = {}  # name -> (label, value, format, provenance, source, judgment) for inputs.json
INP = {}  # name -> cell address on Inputs
INV = {}  # name -> value, so labels and notes can quote the inputs instead of repeating them
row = 5
def inp(name, label, value, fmt=None, prov="", source="", judgment=False, comment=None):
    global row
    ws.cell(row=row, column=1, value=label).font = BLACK
    c = ws.cell(row=row, column=2, value=value)
    c.font = BLUE
    if fmt: c.number_format = fmt
    if judgment: c.fill = YELLOW
    ws.cell(row=row, column=3, value=prov).font = GREY
    ws.cell(row=row, column=5, value=source).font = GREY
    if comment: c.comment = Comment(comment, "model")
    INP[name] = f"Inputs!$B${row}"
    INV[name] = value
    LAB[name] = (label, value, fmt, prov, source, judgment)
    row += 1

def blank(n=1):
    global row
    row += n

header(ws, row, "Market data (as of Oct 15–17, 2026)", 5); row += 1
ws.cell(row=row, column=1, value="Item").font = BOLD; ws.cell(row=row, column=2, value="Value").font = BOLD; ws.cell(row=row, column=3, value="Provenance").font = BOLD; ws.cell(row=row, column=5, value="Source / note").font = BOLD; row += 1
inp("price", "Share price ($)", 32.00, USD2, "MARKET", "Benzinga quote page, last updated Oct 15, 2026 (header shows $31.90; FAQ shows $32.12).")
inp("shares", "Diluted shares (millions)", 529.5, NUM1, "REPORTED", "Q2 2026 earnings release: diluted average shares 529.5M.")
inp("dps", "Dividend per share ($/yr)", 1.85, USD2, "MARKET", "Benzinga: 5.80% yield at $31.90 → $1.85/share; policy maintained through 2026 per Q4 2025 call.")
inp("price_date", "Price date", "Oct 15, 2026", None, "MARKET", "Benzinga quote page.")
inp("wk52_lo", "52-week low ($)", 29, USD2, "MARKET", "Benzinga, Oct 15, 2026 (rounded to the dollar).")
inp("wk52_hi", "52-week high ($)", 50, USD2, "MARKET", "Benzinga, Oct 15, 2026 (rounded to the dollar).")
inp("chg_ytd", "Share price change, year to date", -0.20, PCT, "MARKET", "Benzinga, Oct 15, 2026 (rounded).")
inp("chg_12m", "Share price change, last 12 months", -0.23, PCT, "MARKET", "Benzinga, Oct 15, 2026 (rounded).")
inp("gross_debt", "Gross debt", 9217, USD, "REPORTED", "Q2 2026 balance sheet: notes payable/current maturities $1,002M + long-term debt $8,215M. Excludes $2,131M nonrecourse VIE liabilities.")
inp("cash", "Cash and temporary investments", 726, USD, "REPORTED", "Q2 2026 balance sheet.")
blank()
header(ws, row, "What behaves like debt but isn't in net debt (FY2025 10-K, filed Feb 2026; balances at Dec 31, 2025)", 5); row += 1
inp("pen_us_q_assets", "Pension, U.S. qualified plan: assets", 8482, USD, "REPORTED", "10-K MD&A: 'U.S. qualified pension $ 8,116 $ 8,482' (benefit obligation, fair value of plan assets).")
inp("pen_us_q_pbo", "Pension, U.S. qualified plan: obligation (PBO)", 8116, USD, "REPORTED", "Same table. The plan is overfunded: 'a surplus of $366 million'.")
inp("pen_uk_surplus", "Pension, UK qualified plan (DS Smith): surplus", 112, USD, "REPORTED", "10-K MD&A: 'an overfunded UK qualified pension with a surplus of $112 million.'")
inp("pen_us_nq", "Pension, U.S. nonqualified plans: obligation, no assets", 218, USD, "REPORTED", "10-K MD&A: 'U.S. nonqualified pension 218 —'. Paid as benefits come due ($49M in 2025).")
inp("pen_net", "Pensions, all plans: assets above obligations (the filing's net)", 144, USD, "REPORTED",
    "10-K MD&A: 'the projected benefit obligation ... was approximately $144 million lower than the fair value of plan assets.' (Risk Factors says $148M; the plan table gives 144.)")
inp("pen_surplus", "Pension plans in surplus, total (not counted by default)", 486, USD, "REPORTED", "10-K balance sheet: Overfunded Pension Plan Assets $486M.")
inp("pen_deficit", "Pension plans in deficit, total (counted, after tax)", 342, USD, "REPORTED",
    "Surplus plans $486M less the $144M net = $342M owed by plans in deficit (balance sheet: Underfunded Pension Benefit Obligations $316M long-term, the rest current).")
inp("opeb", "Retiree medical and postemployment benefits (counted, after tax)", 133, USD, "REPORTED", "10-K balance sheet: Postretirement and Postemployment Benefit Obligation $133M (long-term).")
inp("op_leases", "Operating lease liabilities (not counted: the lease cost is already inside EBITDA)", 707, USD, "REPORTED",
    "Dec 31, 2025: $221M current + $486M long-term (Q1 2026 10-Q lease table, December 31, 2025 column; the 10-K balance sheet shows Long-Term Lease Obligations $486M and Right of Use Assets $697M). Under U.S. GAAP operating lease cost is an operating expense, so EBITDA is after rent; adding the liability too would count it twice.")
inp("fin_leases", "Finance lease liabilities (not counted again: already inside gross debt)", 71, USD, "REPORTED", "Dec 31, 2025: $17M in current maturities of debt + $54M in long-term debt (Q1 2026 10-Q lease table, December 31, 2025 column); the 10-K's debt maturities 'include financing lease obligations'.")
inp("timber_dtl", "Deferred tax on the Temple-Inland timber notes, due 2027 (counted)", 487, USD, "REPORTED",
    "10-K, contractual obligations: '$487 million deferred tax liability related to the Temple-Inland timber monetization, which will be settled with the maturity of the notes in 2027.'")
inp("vie_assets", "Timber notes held by the special purpose entities (counted, net of the loans)", 2349, USD, "REPORTED", "10-K balance sheet: Long-Term Financial Assets of Variable Interest Entities $2,349M.")
inp("vie_liab", "Nonrecourse loans against those notes", 2127, USD, "REPORTED",
    "10-K balance sheet: Long-Term Nonrecourse Financial Liabilities of Variable Interest Entities $2,127M. Not in gross debt above; the notes repay them in 2027.")
inp("debt_2027", "Regular debt maturing in 2027 (context, not added)", 2584, USD, "REPORTED", "10-K contractual obligations: debt maturities 2026 $992M, 2027 $2,584M, 2028 $739M.")
inp("utb", "Unrecognized tax benefits (not counted: timing and amount uncertain)", 344, USD, "REPORTED", "10-K contractual obligations footnote: 'approximately $344 million of unrecognized tax benefits'.")
inp("brazil_tax", "Brazil tax matter via the Sylvamo agreement, IP's share (not counted: contingent, not booked)", 274, USD, "REPORTED",
    "10-K Risk Factors: 'our potential liability for such assessments would currently be approximately $274 million' (assessments total about $394M, 2007–2015 tax years).")
inp("count_surplus", "Count the pension surplus too? (1 = yes, 0 = no)", 0, None, "JUDGMENT",
    "Default no: a surplus can't be spent without an excise tax or years of contribution holidays. Set 1 to net it against the deficits.", judgment=True)
inp("pension_other", "Debt-like items, after tax (from the 10-K lines above)", None, USD, "DERIVED",
    "Pension deficits + retiree medical, after tax; minus the surplus if counted; plus the 2027 timber tax; minus the equity in the timber structure (notes less loans). Operating leases excluded.")
blank()

header(ws, row, "Latest reported and guided financials", 5); row += 1
inp("rev_2025_na", "2025 pro forma net sales — Packaging Solutions NA", 15000, USD, "REPORTED", "Q4 2025 call: 'over $15B' pro forma for the North American business (rounded).")
inp("rev_2025_emea", "2025 pro forma net sales — Packaging Solutions EMEA", 8500, USD, "REPORTED", "Q4 2025 call: about $8.5B pro forma for EMEA (DS Smith full year).")
inp("ebitda_2025_na", "2025 pro forma adj. EBITDA — NA", 2300, USD, "REPORTED", "Q4 2025 call: about $2.3B NA.")
inp("ebitda_2025_emea", "2025 pro forma adj. EBITDA — EMEA", 800, USD, "REPORTED", "Q4 2025 call: about $800M EMEA.")
inp("fcf_2025", "2025 free cash flow", -159, USD, "REPORTED", "Q4 2025 call: free cash flow of negative $159M, reflecting transformation investment.")
inp("rev_2026_h1", "H1 2026 net sales", 11975, USD, "REPORTED", "Q2 2026 release: six-month net sales $11,975M.")
inp("ebitda_2026_h1", "H1 2026 adj. EBITDA", 1264, USD, "REPORTED", "Q2 2026 release: six-month adjusted EBITDA $1,264M (Q1 $677M, Q2 $587M).")
inp("capex_2026_h1", "H1 2026 capital expenditures", 1050, USD, "REPORTED", "Q2 2026 release: six-month capex $1,050M.")
inp("ebitda_2026_guide_lo", "2026 adj. EBITDA guidance — low", 3200, USD, "GUIDANCE", "Q2 2026 release: $3.2–3.4B. (January: $3.5–3.7B; cut in April on demand, inflation, Middle East freight.)")
inp("ebitda_2026_guide_hi", "2026 adj. EBITDA guidance — high", 3400, USD, "GUIDANCE", "Q2 2026 release.")
inp("ebitda_2026_guide_jan_lo", "2026 adj. EBITDA guidance in January — low", 3500, USD, "GUIDANCE", "Q4 2025 call (Jan 29, 2026): $3.5–3.7B, cut to $3.2–3.4B in April.")
inp("ebitda_2026_guide_jan_hi", "2026 adj. EBITDA guidance in January — high", 3700, USD, "GUIDANCE", "Q4 2025 call (Jan 29, 2026).")
inp("ebitda_2026_na_guide", "2026 adj. EBITDA guidance — NA (midpoint)", 2425, USD, "GUIDANCE", "April 2026: NA $2.35–2.5B (from $2.5–2.6B).")
inp("ebitda_2026_emea_guide", "2026 adj. EBITDA guidance — EMEA (midpoint)", 950, USD, "GUIDANCE", "April 2026: EMEA $0.9–1.0B (from $1.0–1.1B).")
inp("da_2026", "2026 depreciation & amortization", 1965, USD, "GUIDANCE", "Q2 2026 release outlook reconciliation: D&A $1,965M.")
inp("interest_2026", "2026 interest expense", 370, USD, "GUIDANCE", "Q2 2026 release outlook reconciliation.")
inp("fcf_2026_guide_mid", "2026 free cash flow guidance (Jan, midpoint)", 400, USD, "GUIDANCE", "Q4 2025 call: $300–500M FCF for 2026 (pre-April cut; not updated in the Q2 release).")
inp("sep_costs_2026", "Separation / transaction costs expected 2026–27", 250, USD, "JUDGMENT", "Q2 2026 alone: $43M separation costs + $5M NORPAC costs + $9M 80/20 severance. Annualized run-rate ~$250M through the spin; your call.", judgment=True)
blank()

header(ws, row, "The story, as numbers (Damodaran: narrative → numbers)", 5); row += 1
inp("price_index_ytd", "2026 containerboard index increase recognized YTD ($/ton, linerboard)", 170, USD, "MARKET", "Fastmarkets RISI via Packaging Dive (Sept 18, 2026): −20 Feb, +40 Mar, +30 Apr, +50 Jun, +70 Sep ≈ +$170/ton linerboard; medium +$100 in Sept alone.")
inp("price_sens", "EBITDA per $1/ton of price (annualized, NA)", 9, USD, "GUIDANCE", "CFO Lance Loeffler, Q4 2025 call: each $10/ton ≈ $90M annualized adjusted EBITDA.")
inp("price_in_2026", "Price benefit already inside 2026 NA guidance", 175, USD, "GUIDANCE", "Q1 2026 call: IP expects pricing to add ~$175M to NA in 2026, mostly in H2.")
inp("stick_pct", "Share of the YTD index gain that STICKS through 2027 (floor judgment)", 0.50, PCT, "JUDGMENT", "Q1 for Jason: of the $170/ton announced, what holds at the box level in 12 months? Green Markets found no evidence the third round reached the open market (Sept). Default 50%.", judgment=True)
inp("vol_2027_na", "NA box volume growth 2027", 0.01, PCT, "JUDGMENT", "Market flat in 2026 (TTOBMA cut containerboard to −2%); IP claims +3–4 pts vs market. Default +1%.", judgment=True)
inp("contrib_margin", "Contribution margin on incremental NA volume", 0.30, PCT, "JUDGMENT", "Incremental box volume through existing mills/plants; typical converter contribution 25–35%.", judgment=True)
inp("costout_2027", "Incremental 80/20 cost-out 2027 (NA + EMEA)", 300, USD, "JUDGMENT", "~$510M run-rate benefits in 2025, 'similar' expected in 2026 (Q4 2025 call); EMEA >$200M run-rate from 31 closures. Default $300M for 2027, fading.", judgment=True)
inp("inflation_2027", "Cost inflation 2027 (OCC, energy, labor, freight)", 250, USD, "JUDGMENT", "OCC back to ~$80–100/ton from 2025 lows (5-yr avg $140); freight up on Middle East; EMEA inflation ~$100M in 2026 guide. Default $250M.", judgment=True)
inp("emea_2027_delta", "EMEA EBITDA change 2027 (box-price catch-up + cost-out, net)", 150, USD, "JUDGMENT", "Q2 2026: EMEA lost money because paper prices rose faster than box prices. Catch-up + cost-out vs soft demand. Default +$150M.", judgment=True)
inp("midcycle_na_margin", "Mid-cycle NA adj. EBITDA margin (2028+)", 0.17, PCT, "JUDGMENT", "SW 15% LTM, PKG 21% LTM, IP NA 2026E ~16.5%. Bull = toward PKG after 80/20; bear = SW. Default 17%.", judgment=True)
inp("midcycle_emea_margin", "Mid-cycle EMEA adj. EBITDA margin (2028+)", 0.11, PCT, "JUDGMENT", "DS Smith historically ~12–13%; 2026E ~10%. Default 11%.", judgment=True)
inp("rev_growth_lt", "Revenue growth 2028–2030 (nominal)", 0.025, PCT, "JUDGMENT", "Box demand ≈ GDP-ish plus price; 2.5% nominal.", judgment=True)
inp("capex_2027", "Capex 2027", 1900, USD, "JUDGMENT", "2026 run-rate ~$2.1B incl. Riverdale conversion ($250M project) and $225M Mississippi greenfield (opens Q4 2027). Default $1.9B.", judgment=True)
inp("capex_pct_lt", "Capex as % of revenue 2028+ (maintenance + modest growth)", 0.065, PCT, "JUDGMENT", "Paper is capex-heavy: D&A ≈ 8% of sales today. 6.5% assumes the transformation capex rolls off. Your 'assets in the ground are expensive to run' point lives here.", judgment=True)
inp("da_pct_lt", "D&A as % of revenue 2027+", 0.080, PCT, "JUDGMENT", "2026: $1,965M on ~$24B = 8.2%.", judgment=True)
inp("nwc_pct", "Net working capital investment, % of incremental revenue", 0.10, PCT, "JUDGMENT", "Packaging converters typically carry 8–12% of sales in NWC.", judgment=True)
inp("tax_rate", "Cash tax rate", 0.24, PCT, "JUDGMENT", "US 21% federal + state, blended with EMEA. Check against the 10-K effective rate.", judgment=True)
blank()

header(ws, row, "Discount rate (WACC build)", 5); row += 1
inp("rf", "Risk-free rate (10-yr US Treasury)", 0.040, PCT, "MARKET", "Placeholder near the Oct 2026 10-yr yield — refresh before publishing.", judgment=True)
inp("beta", "Equity beta", 1.10, '0.00', "JUDGMENT", "IP 5-yr regression beta has run ~1.0–1.2; Damodaran paper/forest products industry beta ~1.0. Default 1.10.", judgment=True)
inp("erp", "Equity risk premium", 0.045, PCT, "JUDGMENT", "Damodaran implied ERP has run 4–5% in 2025–26.", judgment=True)
inp("kd", "Pre-tax cost of debt", 0.055, PCT, "MARKET", "IP senior unsecured notes have traded ~5–6% yield (BBB-range credit). Refresh.", judgment=True)
inp("g", "Terminal growth rate", 0.020, PCT, "JUDGMENT", "At or below long-run nominal GDP for a mature, capex-heavy business.", judgment=True)
inp("exit_mult", "Exit EV/EBITDA multiple (cross-check)", 8.0, MULT, "JUDGMENT", "IP's own 10-yr average ~8–9x; SW 7.2x and PKG 12.2x today.", judgment=True)
blank()

header(ws, row, "Peers (Multiples.vc, Oct 15–17, 2026)", 5); row += 1
inp("pkg_ev_ebitda", "PKG EV/EBITDA (LTM)", 12.2, MULT, "MARKET", "Multiples.vc Oct 17, 2026; EV $25B, LTM revenue $9.8B, EBITDA margin 21%.")
inp("pkg_margin", "PKG EBITDA margin (LTM)", 0.21, PCT, "MARKET", "Multiples.vc.")
inp("sw_ev_ebitda", "Smurfit Westrock EV/EBITDA (LTM)", 7.2, MULT, "MARKET", "Multiples.vc Oct 15, 2026; EV $35B, LTM revenue $32B, EBITDA margin 15%.")
inp("sw_margin", "Smurfit Westrock EBITDA margin (LTM)", 0.15, PCT, "MARKET", "Multiples.vc.")
blank()

header(ws, row, "Scenario switch", 5); row += 1
inp("scenario", "Active scenario (1 = Bear, 2 = Mid-cycle, 3 = Bull) — drives the Scenarios sheet overrides", 2, '0', "JUDGMENT", "The Model sheet uses the Inputs above directly; the Scenarios sheet runs all three side by side.", judgment=True)
# the debt-like total is a formula over the 10-K lines, written once the tax rate's cell exists
c = ws.cell(row=int(INP["pension_other"].split("$")[-1]), column=2,
            value=f"=({INP['pen_deficit']}+{INP['opeb']})*(1-{INP['tax_rate']})-{INP['pen_surplus']}*(1-{INP['tax_rate']})*{INP['count_surplus']}"
                  f"+{INP['timber_dtl']}-({INP['vie_assets']}-{INP['vie_liab']})")
c.font = BLACK; c.number_format = USD
# notes that quote other inputs are rewritten once every input exists, so they can't go stale
_rev_na_26 = INV["rev_2026_h1"] * 2 * INV["rev_2025_na"] / (INV["rev_2025_na"] + INV["rev_2025_emea"])
_note = (f"SW {INV['sw_margin']:.0%} LTM, PKG {INV['pkg_margin']:.0%} LTM, IP NA 2026E ~{INV['ebitda_2026_na_guide'] / _rev_na_26:.1%}. "
         f"Bull = toward PKG after 80/20; bear = below SW. Default {INV['midcycle_na_margin']:.0%}.")
ws.cell(row=int(INP["midcycle_na_margin"].split("$")[-1]), column=5, value=_note).font = GREY
LAB["midcycle_na_margin"] = LAB["midcycle_na_margin"][:4] + (_note,) + LAB["midcycle_na_margin"][5:]

# =====================================================================================
# MODEL  (2025A pro forma, 2026E guidance, 2027E bridge, 2028–2030E mid-cycle)
# =====================================================================================
m = wb.create_sheet("Model")
setw(m, [44, 13, 13, 13, 13, 13, 13, 44])
m["A1"] = "Projection model — enterprise (NA + EMEA retained; see SOTP for the separation)"; m["A1"].font = H1
m["A2"] = "$ millions. 2026E = guidance midpoint. 2027E = driver bridge from the story. 2028E+ = converge to mid-cycle margins."; m["A2"].font = GREY
years = ["2025A", "2026E", "2027E", "2028E", "2029E", "2030E"]
cols = ["B", "C", "D", "E", "F", "G"]
r = 4
m.cell(row=r, column=1, value="Fiscal year").font = BOLD
for c, y in zip(cols, years):
    cell = m[f"{c}{r}"]; cell.value = y; cell.font = BOLD; cell.alignment = Alignment(horizontal="right")
m[f"H{r}"] = "How it's built"; m[f"H{r}"].font = BOLD
R = {}
def mrow(key, label, formulas, fmt=USD, note="", bold=False):
    global r
    r += 1
    m.cell(row=r, column=1, value=label).font = BOLD if bold else BLACK
    for c, f in zip(cols, formulas):
        if f is None: continue
        cell = m[f"{c}{r}"]; cell.value = f; cell.number_format = fmt
        cell.font = GREEN if (isinstance(f, str) and "Inputs!" in f and f.count("Inputs!") == 1 and f.startswith("=Inputs!")) else BLACK
        if bold: cell.font = Font(name=FONT, bold=True, size=10, color=cell.font.color.rgb if cell.font.color else "000000")
    if note: m.cell(row=r, column=8, value=note).font = GREY
    R[key] = r

def I(name): return INP[name]

sub(m, r + 1, "Revenue", 8); r += 1
mrow("rev_na", "Net sales — NA",
     [f"={I('rev_2025_na')}",
      f"=B{{rev_na}}*(1+0.0)+{I('price_in_2026')}/0.30",   # placeholder replaced below
      None, None, None, None], note="")
# rebuild rev_na with proper formulas after we know the row
rr = R["rev_na"]
m[f"C{rr}"] = f"={I('rev_2026_h1')}*2*{I('rev_2025_na')}/({I('rev_2025_na')}+{I('rev_2025_emea')})"
m[f"C{rr}"].number_format = USD
m[f"D{rr}"] = f"=C{rr}*(1+{I('vol_2027_na')})+{I('stick_pct')}*{I('price_index_ytd')}*{I('price_sens')}-{I('price_in_2026')}"
m[f"D{rr}"].number_format = USD
for c, p in zip(["E", "F", "G"], ["D", "E", "F"]):
    m[f"{c}{rr}"] = f"={p}{rr}*(1+{I('rev_growth_lt')})"; m[f"{c}{rr}"].number_format = USD
m[f"H{rr}"] = "2026E: H1×2 split by 2025 mix. 2027E: volume growth + the incremental price that sticks (price $ passes through revenue 1:1). 2028+: long-term growth."; m[f"H{rr}"].font = GREY

mrow("rev_emea", "Net sales — EMEA",
     [f"={I('rev_2025_emea')}",
      f"={I('rev_2026_h1')}*2*{I('rev_2025_emea')}/({I('rev_2025_na')}+{I('rev_2025_emea')})",
      f"=C{{rev_emea}}*(1+{I('rev_growth_lt')})", None, None, None])
rr = R["rev_emea"]
m[f"D{rr}"] = f"=C{rr}*(1+{I('rev_growth_lt')})"; m[f"D{rr}"].number_format = USD
for c, p in zip(["E", "F", "G"], ["D", "E", "F"]):
    m[f"{c}{rr}"] = f"={p}{rr}*(1+{I('rev_growth_lt')})"; m[f"{c}{rr}"].number_format = USD
m[f"H{rr}"] = "EMEA grows with the long-term rate; demand soft, price catch-up handled in EBITDA."; m[f"H{rr}"].font = GREY

mrow("rev", "Net sales — enterprise", [f"={c}{R['rev_na']}+{c}{R['rev_emea']}" for c in cols], bold=True)
mrow("rev_g", "  growth", [None] + [f"={c}{R['rev']}/{p}{R['rev']}-1" for c, p in zip(cols[1:], cols[:-1])], fmt=PCT)

sub(m, r + 1, "Adjusted EBITDA — the bridge", 8); r += 1
mrow("ebitda_na", "Adj. EBITDA — NA",
     [f"={I('ebitda_2025_na')}", f"={I('ebitda_2026_na_guide')}", None, None, None, None])
rr = R["ebitda_na"]
m[f"D{rr}"] = (f"=C{rr}+({I('stick_pct')}*{I('price_index_ytd')}*{I('price_sens')}-{I('price_in_2026')})"
               f"+C{R['rev_na']}*{I('vol_2027_na')}*{I('contrib_margin')}"
               f"+{I('costout_2027')}*C{R['rev_na']}/C{R['rev']}-{I('inflation_2027')}*C{R['rev_na']}/C{R['rev']}")
m[f"D{rr}"].number_format = USD
for c in ["E", "F", "G"]:
    m[f"{c}{rr}"] = f"={c}{R['rev_na']}*{I('midcycle_na_margin')}"; m[f"{c}{rr}"].number_format = USD
m[f"H{rr}"] = "2027E = 2026 guide + price that sticks (net of what's already in) + volume × contribution + cost-out − inflation (NA share of both). 2028+ = mid-cycle margin × revenue — DELIBERATELY below 2027: a price-increase year is a peak, and a cyclical is valued at mid-cycle, not at the peak (Damodaran)."; m[f"H{rr}"].font = GREY

mrow("ebitda_emea", "Adj. EBITDA — EMEA",
     [f"={I('ebitda_2025_emea')}", f"={I('ebitda_2026_emea_guide')}", None, None, None, None])
rr = R["ebitda_emea"]
m[f"D{rr}"] = f"=C{rr}+{I('emea_2027_delta')}+({I('costout_2027')}-{I('inflation_2027')})*C{R['rev_emea']}/C{R['rev']}"
m[f"D{rr}"].number_format = USD
for c in ["E", "F", "G"]:
    m[f"{c}{rr}"] = f"={c}{R['rev_emea']}*{I('midcycle_emea_margin')}"; m[f"{c}{rr}"].number_format = USD
m[f"H{rr}"] = "2027E = 2026 guide + box-price catch-up + EMEA share of cost-out − inflation. 2028+ = mid-cycle margin."; m[f"H{rr}"].font = GREY

mrow("corp", "Corporate / other",
     [f"=-{I('ebitda_2025_na')}*0+{I('ebitda_2025_emea')}*0-100",
      f"=({I('ebitda_2026_guide_lo')}+{I('ebitda_2026_guide_hi')})/2-{I('ebitda_2026_na_guide')}-{I('ebitda_2026_emea_guide')}",
      f"=C{{corp}}", f"=C{{corp}}", f"=C{{corp}}", f"=C{{corp}}"])
rr = R["corp"]
m[f"B{rr}"] = -100; m[f"B{rr}"].font = BLUE; m[f"B{rr}"].number_format = USD
for c in ["D", "E", "F", "G"]:
    m[f"{c}{rr}"] = f"=C{rr}"; m[f"{c}{rr}"].number_format = USD
m[f"H{rr}"] = "2026E = enterprise guide midpoint − segment guides (the unallocated corporate drag); held flat. 2025A −$100M is an estimate (blue)."; m[f"H{rr}"].font = GREY

mrow("ebitda", "Adj. EBITDA — enterprise", [f"={c}{R['ebitda_na']}+{c}{R['ebitda_emea']}+{c}{R['corp']}" for c in cols], bold=True)
mrow("ebitda_m", "  margin", [f"={c}{R['ebitda']}/{c}{R['rev']}" for c in cols], fmt=PCT)
mrow("ebitda_na_m", "  NA margin", [f"={c}{R['ebitda_na']}/{c}{R['rev_na']}" for c in cols], fmt=PCT, note="Compare: PKG 21%, SW 15%.")
mrow("ebitda_emea_m", "  EMEA margin", [f"={c}{R['ebitda_emea']}/{c}{R['rev_emea']}" for c in cols], fmt=PCT)

sub(m, r + 1, "From EBITDA to unlevered free cash flow", 8); r += 1
mrow("da", "Depreciation & amortization",
     [f"=-{I('da_2026')}*0.95", f"=-{I('da_2026')}"] + [f"=-{c}{R['rev']}*{I('da_pct_lt')}" for c in cols[2:]],
     note="2026E = guidance; 2027+ = % of revenue. 2025A approximated at 95% of 2026.")
mrow("ebit", "EBIT (adjusted)", [f"={c}{R['ebitda']}+{c}{R['da']}" for c in cols], bold=True)
mrow("tax", "Cash taxes on EBIT", [f"=-MAX(0,{c}{R['ebit']})*{I('tax_rate')}" for c in cols], note="No tax benefit modeled on an EBIT loss (conservative).")
mrow("nopat", "NOPAT", [f"={c}{R['ebit']}+{c}{R['tax']}" for c in cols])
mrow("da_add", "Add back D&A", [f"=-{c}{R['da']}" for c in cols])
mrow("capex", "Capital expenditures",
     [-2000, f"=-{I('capex_2026_h1')}*2", f"=-{I('capex_2027')}"] + [f"=-{c}{R['rev']}*{I('capex_pct_lt')}" for c in cols[3:]],
     note="2026E = H1 × 2. 2027E = input. 2028+ = % of revenue. 2025A ~$2.0B estimate.")
mrow("nwc", "Investment in net working capital", [None] + [f"=-({c}{R['rev']}-{p}{R['rev']})*{I('nwc_pct')}" for c, p in zip(cols[1:], cols[:-1])])
mrow("sep", "Separation / transaction costs (cash)", [None, f"=-{I('sep_costs_2026')}", f"=-{I('sep_costs_2026')}*0.5", 0, 0, 0], note="Below EBITDA because 'adjusted' excludes them; the cash still leaves. 2027 at half, then done.")
mrow("ufcf", "Unlevered free cash flow", [f"={c}{R['nopat']}+{c}{R['da_add']}+{c}{R['capex']}+IF(ISNUMBER({c}{R['nwc']}),{c}{R['nwc']},0)+IF(ISNUMBER({c}{R['sep']}),{c}{R['sep']},0)" for c in cols], bold=True)
mrow("ufcf_m", "  UFCF / revenue", [f"={c}{R['ufcf']}/{c}{R['rev']}" for c in cols], fmt=PCT)

sub(m, r + 1, "Sanity checks against what the company says", 8); r += 1
mrow("chk_fcf", "Levered FCF proxy = UFCF − after-tax interest", [None, f"=C{R['ufcf']}-{I('interest_2026')}*(1-{I('tax_rate')})", f"=D{R['ufcf']}-{I('interest_2026')}*(1-{I('tax_rate')})", None, None, None], note="Compare 2026E to the $300–500M January FCF guide and 2025's −$159M actual.")
mrow("chk_div", "Dividends paid", [f"=-{I('dps')}*{I('shares')}"] * 6, note=f"${INV['dps']:.2f} × {INV['shares']}M ≈ ${INV['dps'] * INV['shares']:,.0f}M/yr. Covered only when the row above exceeds this.")
mrow("chk_2025", "2025A reported free cash flow (actual)", [f"={I('fcf_2025')}", None, None, None, None, None], note="The 2025A model column is a pro forma approximation (no DS Smith deal costs, severance cash, pension, or NWC swing) and will not tie to this −$159M; it is there for scale, not for valuation. Nothing is discounted from 2025 or 2026.")
mrow("chk_cov", "  dividend coverage by levered FCF", [None, f"=-C{R['chk_fcf']}/C{R['chk_div']}", f"=-D{R['chk_fcf']}/D{R['chk_div']}", None, None, None], fmt=MULT)

for rr_ in range(5, r + 1):
    for c in ["B", "C", "D", "E", "F", "G"]:
        m[f"{c}{rr_}"].alignment = Alignment(horizontal="right")

# =====================================================================================
# DCF
# =====================================================================================
d = wb.create_sheet("DCF")
setw(d, [44, 13, 13, 13, 13, 13, 60])
d["A1"] = "Discounted cash flow — enterprise, retained EMEA"; d["A1"].font = H1
d["A2"] = "Valuation date Oct 17, 2026. Cash flows discounted from mid-2027 onward; 2026E cash flow excluded (mostly in the past and already in net debt by year-end)."; d["A2"].font = GREY
r = 4
header(d, r, "WACC", 7); r += 1
D = {}
def drow(key, label, f, fmt=USD, note="", bold=False, col="B"):
    global r
    d.cell(row=r, column=1, value=label).font = BOLD if bold else BLACK
    cell = d[f"{col}{r}"]; cell.value = f; cell.number_format = fmt
    cell.font = GREEN if (isinstance(f, str) and f.startswith("=Inputs!")) else (BOLD if bold else BLACK)
    if note: d.cell(row=r, column=7, value=note).font = GREY
    D[key] = f"{col}{r}"; r += 1

drow("mcap", "Market capitalization", f"={I('price')}*{I('shares')}")
drow("netdebt", "Net debt (gross debt − cash)", f"={I('gross_debt')}-{I('cash')}")
drow("debtlike", "Debt-like items from the 10-K, after tax (Inputs)", f"={I('pension_other')}", note="Pension deficits and retiree medical after tax, the 2027 timber tax net of the notes' equity. Off the equity, not in the WACC weights: these aren't financing.")
drow("we", "Weight of equity (market)", f"={D['mcap']}/({D['mcap']}+{D['netdebt']})", PCT)
drow("wd", "Weight of debt (market)", f"=1-{D['we']}", PCT)
drow("ke", "Cost of equity = rf + β × ERP", f"={I('rf')}+{I('beta')}*{I('erp')}", PCT)
drow("kd_at", "After-tax cost of debt", f"={I('kd')}*(1-{I('tax_rate')})", PCT)
drow("wacc", "WACC", f"={D['we']}*{D['ke']}+{D['wd']}*{D['kd_at']}", PCT, bold=True, note="Damodaran would push back on using today's depressed equity weight for a target structure; a 30/70 D/E target gives a similar number here.")
r += 1
header(d, r, "Cash flows", 7); r += 1
d.cell(row=r, column=1, value="Fiscal year").font = BOLD
dcols = ["B", "C", "D", "E"]; dyears = ["2027E", "2028E", "2029E", "2030E"]; mcols = ["D", "E", "F", "G"]
for c, y in zip(dcols, dyears):
    d[f"{c}{r}"] = y; d[f"{c}{r}"].font = BOLD; d[f"{c}{r}"].alignment = Alignment(horizontal="right")
hdr = r; r += 1
d.cell(row=r, column=1, value="Unlevered free cash flow").font = BLACK
for c, mc in zip(dcols, mcols):
    d[f"{c}{r}"] = f"=Model!{mc}{R['ufcf']}"; d[f"{c}{r}"].font = GREEN; d[f"{c}{r}"].number_format = USD
ufcf_r = r; r += 1
d.cell(row=r, column=1, value="Years from valuation date (mid-year)").font = BLACK
for i, c in enumerate(dcols):
    d[f"{c}{r}"] = 0.7 + i; d[f"{c}{r}"].font = BLUE; d[f"{c}{r}"].number_format = '0.0'
d["G" + str(r)] = "Oct 2026 → mid-2027 ≈ 0.7 years; then annual."; d[f"G{r}"].font = GREY
t_r = r; r += 1
d.cell(row=r, column=1, value="Discount factor").font = BLACK
for c in dcols:
    d[f"{c}{r}"] = f"=1/(1+{D['wacc']})^{c}{t_r}"; d[f"{c}{r}"].number_format = '0.000'
df_r = r; r += 1
d.cell(row=r, column=1, value="Present value of UFCF").font = BLACK
for c in dcols:
    d[f"{c}{r}"] = f"={c}{ufcf_r}*{c}{df_r}"; d[f"{c}{r}"].number_format = USD
pv_r = r; r += 2

header(d, r, "Terminal value — two ways", 7); r += 1
drow("tv_gg", "Gordon growth: UFCF₂₀₃₀ × (1+g) / (WACC − g)", f"=E{ufcf_r}*(1+{I('g')})/({D['wacc']}-{I('g')})", note="Requires WACC > g; the model errors loudly otherwise, by design.")
drow("tv_em", "Exit multiple: EBITDA₂₀₃₀ × multiple", f"=Model!G{R['ebitda']}*{I('exit_mult')}")
drow("tv_gg_mult", "  implied exit multiple of the Gordon TV", f"={D['tv_gg']}/Model!G{R['ebitda']}", MULT, note="If this is far from 7–12x, the g/WACC pair is telling a story the comps don't.")
drow("tv_em_g", "  implied perpetual growth of the exit-multiple TV", f"=({D['tv_em']}*{D['wacc']}-E{ufcf_r})/({D['tv_em']}+E{ufcf_r})", PCT)
drow("tv", "Terminal value used (average of the two)", f"=AVERAGE({D['tv_gg']},{D['tv_em']})", bold=True)
drow("pv_tv", "PV of terminal value", f"={D['tv']}*E{df_r}")
r += 1
header(d, r, "Enterprise value → equity value → per share", 7); r += 1
drow("pv_sum", "Sum of PV of 2027–2030 UFCF", f"=SUM(B{pv_r}:E{pv_r})")
drow("ev", "Enterprise value", f"={D['pv_sum']}+{D['pv_tv']}", bold=True)
drow("tv_share", "  terminal value as % of EV", f"={D['pv_tv']}/{D['ev']}", PCT, note="Damodaran's honesty check: for a mature company 60–80% is normal; it means the answer IS the terminal assumptions.")
drow("eq", "Equity value = EV − net debt − debt-like", f"={D['ev']}-{D['netdebt']}-{D['debtlike']}", bold=True)
drow("ps", "Value per share ($)", f"={D['eq']}/{I('shares')}", USD2, bold=True)
drow("upside", "  vs. market price", f"={D['ps']}/{I('price')}-1", PCT)
drow("impl_mult", "  implied EV / 2027E EBITDA", f"={D['ev']}/Model!D{R['ebitda']}", MULT)
drow("mkt_mult", "  market EV / 2027E EBITDA (today)", f"=({D['mcap']}+{D['netdebt']})/Model!D{R['ebitda']}", MULT)
drow("mkt_mult26", "  market EV / 2026E EBITDA (today)", f"=({D['mcap']}+{D['netdebt']})/Model!C{R['ebitda']}", MULT, note="PKG 12.2x and SW 7.2x LTM for comparison.")

# =====================================================================================
# SCENARIOS (three columns, same chain, compact)
# =====================================================================================
s = wb.create_sheet("Scenarios")
setw(s, [46, 14, 14, 14, 60])
s["A1"] = "Three stories, three numbers"; s["A1"].font = H1
s["A2"] = "Each column is a complete narrative → numbers chain. Blue cells are the scenario levers; everything else is formula. 2027E bridge + mid-cycle 2028–2030 + terminal, same mechanics as the DCF sheet."; s["A2"].font = GREY
r = 4
s.cell(row=r, column=1, value="Lever").font = BOLD
for c, nm in zip(["B", "C", "D"], ["Bear", "Mid-cycle", "Bull"]):
    s[f"{c}{r}"] = nm; s[f"{c}{r}"].font = BOLD; s[f"{c}{r}"].alignment = Alignment(horizontal="right")
s["E4"] = "The story"; s["E4"].font = BOLD
r += 1
S = {}
def srow(key, label, vals, fmt, note="", inputs=True, bold=False):
    global r
    s.cell(row=r, column=1, value=label).font = BOLD if bold else BLACK
    for c, v in zip(["B", "C", "D"], vals):
        cell = s[f"{c}{r}"]; cell.value = v; cell.number_format = fmt
        cell.font = BLUE if inputs else (BOLD if bold else BLACK)
        if inputs: cell.fill = YELLOW
    if note: s.cell(row=r, column=5, value=note).font = GREY
    S[key] = r; r += 1

srow("stick", f"Price that sticks through 2027 (% of ${INV['price_index_ytd']}/ton)", [0.25, 0.50, 0.80], PCT, "Bear: demand −2% and the third round never reaches the box. Bull: capacity cut is real and index-linked contracts catch up.")
srow("vol", "NA volume growth 2027", [-0.02, 0.01, 0.03], PCT, "Bear: TTOBMA's −2%. Bull: IP's '3–4 pts above market'.")
srow("costout", "Cost-out 2027", [200, 300, 450], USD, "Bull assumes 80/20 keeps finding $500M/yr as it did in 2025.")
srow("infl", "Inflation 2027", [350, 250, 200], USD, "Bear: OCC back to its $140 five-year average, freight stays up.")
srow("emea", "EMEA 2027 change", [0, 150, 300], USD, "Bear: box prices never catch paper prices; sale happens at a distressed multiple.")
srow("na_m", "Mid-cycle NA margin 2028+", [0.13, 0.17, 0.20], PCT, "Bear = below SW today; Mid = between SW and PKG; Bull = approaching PKG.")
srow("emea_m", "Mid-cycle EMEA margin 2028+", [0.09, 0.11, 0.13], PCT, "DS Smith's historical range.")
srow("capex", "Capex % of revenue 2028+", [0.075, 0.065, 0.060], PCT, "Bear: the old mills keep eating capital ('assets in the ground').")
srow("wacc", "WACC", [f"=DCF!{D['wacc']}+0.005", f"=DCF!{D['wacc']}", f"=DCF!{D['wacc']}-0.005"], PCT, "±50 bp.")
srow("mult", "Exit multiple", [7.0, 8.0, 9.5], MULT, "Bear = SW today; Bull = halfway to PKG.")
r += 1
sub(s, r, "Derived (formulas)", 5); r += 1
# shared references
rev_na_26 = f"Model!C{R['rev_na']}"; rev_emea_26 = f"Model!C{R['rev_emea']}"; rev_26 = f"Model!C{R['rev']}"
na_26 = f"Model!C{R['ebitda_na']}"; emea_26 = f"Model!C{R['ebitda_emea']}"; corp = f"Model!C{R['corp']}"
def srow_f(key, label, fn, fmt, note="", bold=False):
    srow(key, label, [fn(c) for c in ["B", "C", "D"]], fmt, note, inputs=False, bold=bold)

srow_f("rev_na27", "NA revenue 2027E", lambda c: f"={rev_na_26}*(1+{c}{S['vol']})+{c}{S['stick']}*{I('price_index_ytd')}*{I('price_sens')}-{I('price_in_2026')}", USD)
srow_f("rev_emea27", "EMEA revenue 2027E", lambda c: f"={rev_emea_26}*(1+{I('rev_growth_lt')})", USD)
srow_f("na27", "NA EBITDA 2027E", lambda c: f"={na_26}+({c}{S['stick']}*{I('price_index_ytd')}*{I('price_sens')}-{I('price_in_2026')})+{rev_na_26}*{c}{S['vol']}*{I('contrib_margin')}+({c}{S['costout']}-{c}{S['infl']})*{rev_na_26}/{rev_26}", USD)
srow_f("emea27", "EMEA EBITDA 2027E", lambda c: f"={emea_26}+{c}{S['emea']}+({c}{S['costout']}-{c}{S['infl']})*{rev_emea_26}/{rev_26}", USD)
srow_f("ebitda27", "Enterprise EBITDA 2027E", lambda c: f"={c}{S['na27']}+{c}{S['emea27']}+{corp}", USD, bold=True)
srow_f("rev30", "Enterprise revenue 2030E", lambda c: f"=({c}{S['rev_na27']}+{c}{S['rev_emea27']})*(1+{I('rev_growth_lt')})^3", USD)
srow_f("ebitda30", "Enterprise EBITDA 2030E (mid-cycle)", lambda c: f"={c}{S['rev_na27']}*(1+{I('rev_growth_lt')})^3*{c}{S['na_m']}+{c}{S['rev_emea27']}*(1+{I('rev_growth_lt')})^3*{c}{S['emea_m']}+{corp}", USD, bold=True)
srow_f("ufcf30", "UFCF 2030E", lambda c: f"=({c}{S['ebitda30']}-{c}{S['rev30']}*{I('da_pct_lt')})*(1-{I('tax_rate')})+{c}{S['rev30']}*{I('da_pct_lt')}-{c}{S['rev30']}*{c}{S['capex']}-{c}{S['rev30']}*{I('rev_growth_lt')}/(1+{I('rev_growth_lt')})*{I('nwc_pct')}", USD, "Steady-state: NOPAT + D&A − capex − NWC on one year's growth.")
srow_f("ufcf27", "UFCF 2027E (approx.)", lambda c: f"=({c}{S['ebitda27']}-({c}{S['rev_na27']}+{c}{S['rev_emea27']})*{I('da_pct_lt')})*(1-{I('tax_rate')})+({c}{S['rev_na27']}+{c}{S['rev_emea27']})*{I('da_pct_lt')}-{I('capex_2027')}-{I('sep_costs_2026')}*0.5", USD, "2027 capex and half the separation cost from Inputs; NWC ignored for the approximation.")
srow_f("pv_interim", "PV of 2027–2030 UFCF (linear ramp 2027→2030)", lambda c: f"={c}{S['ufcf27']}/(1+{c}{S['wacc']})^0.7+({c}{S['ufcf27']}*2/3+{c}{S['ufcf30']}/3)/(1+{c}{S['wacc']})^1.7+({c}{S['ufcf27']}/3+{c}{S['ufcf30']}*2/3)/(1+{c}{S['wacc']})^2.7+{c}{S['ufcf30']}/(1+{c}{S['wacc']})^3.7", USD, "Interpolates 2028–29 between the bridge year and the steady state.")
srow_f("tv", "Terminal value (avg of Gordon and exit multiple)", lambda c: f"=AVERAGE({c}{S['ufcf30']}*(1+{I('g')})/({c}{S['wacc']}-{I('g')}),{c}{S['ebitda30']}*{c}{S['mult']})", USD)
srow_f("ev", "Enterprise value", lambda c: f"={c}{S['pv_interim']}+{c}{S['tv']}/(1+{c}{S['wacc']})^3.7", USD, bold=True)
srow_f("eq", "Equity value", lambda c: f"={c}{S['ev']}-DCF!{D['netdebt']}-DCF!{D['debtlike']}", USD)
srow_f("ps", "Value per share ($)", lambda c: f"={c}{S['eq']}/{I('shares')}", USD2, bold=True)
srow_f("up", f"vs. ${INV['price']:.0f} market", lambda c: f"={c}{S['ps']}/{I('price')}-1", PCT)
srow_f("tvs", "Terminal value share of EV", lambda c: f"={c}{S['tv']}/(1+{c}{S['wacc']})^3.7/{c}{S['ev']}", PCT)
srow_f("impl", "Implied EV / 2027E EBITDA", lambda c: f"={c}{S['ev']}/{c}{S['ebitda27']}", MULT)
r += 1
s.cell(row=r, column=1, value="Probability weight (your call)").font = BLACK
for c, v in zip(["B", "C", "D"], [0.30, 0.45, 0.25]):
    s[f"{c}{r}"] = v; s[f"{c}{r}"].font = BLUE; s[f"{c}{r}"].fill = YELLOW; s[f"{c}{r}"].number_format = PCT
S["w"] = r; r += 1
s.cell(row=r, column=1, value="Probability-weighted value per share ($)").font = BOLD
s[f"B{r}"] = f"=SUMPRODUCT(B{S['ps']}:D{S['ps']},B{S['w']}:D{S['w']})/SUM(B{S['w']}:D{S['w']})"; s[f"B{r}"].number_format = USD2; s[f"B{r}"].font = BOLD
s[f"E{r}"] = "Not a price target. It's the expected value of three stories you've written down and can defend line by line."; s[f"E{r}"].font = GREY
S["pw"] = r

# =====================================================================================
# SENSITIVITY (two grids, all formulas)
# =====================================================================================
v = wb.create_sheet("Sensitivity")
setw(v, [40, 12, 12, 12, 12, 12, 12])
v["A1"] = "What actually moves the number"; v["A1"].font = H1
v["A2"] = "Grid 1: value per share vs mid-cycle NA margin (rows) × exit multiple (columns), mid-cycle scenario mechanics. Grid 2: 2027E EBITDA vs how much of the $" + str(INV["price_index_ytd"]) + "/ton sticks."; v["A2"].font = GREY
r = 4
v.cell(row=r, column=1, value="Value / share: NA mid-cycle margin ↓  ×  exit multiple →").font = BOLD
mults = [6.5, 7.5, 8.5, 9.5, 11.0, 12.0]
margins = [0.13, 0.15, 0.17, 0.19, 0.21]
r += 1
for j, mu in enumerate(mults):
    cell = v.cell(row=r, column=2 + j, value=mu); cell.font = BLUE; cell.number_format = MULT; cell.alignment = Alignment(horizontal="right")
hdr_r = r; r += 1
# mid-cycle chain parameterized by margin (row) and multiple (col), using Scenarios mid column for everything else
mid = "C"
for i, mg in enumerate(margins):
    cell = v.cell(row=r, column=1, value=mg); cell.font = BLUE; cell.number_format = PCT; cell.alignment = Alignment(horizontal="right")
    for j in range(len(mults)):
        col = get_column_letter(2 + j)
        ebitda30 = f"(Scenarios!{mid}{S['rev_na27']}*(1+{I('rev_growth_lt')})^3*$A{r}+Scenarios!{mid}{S['rev_emea27']}*(1+{I('rev_growth_lt')})^3*Scenarios!{mid}{S['emea_m']}+{corp})"
        rev30 = f"Scenarios!{mid}{S['rev30']}"
        ufcf30 = f"(({ebitda30}-{rev30}*{I('da_pct_lt')})*(1-{I('tax_rate')})+{rev30}*{I('da_pct_lt')}-{rev30}*Scenarios!{mid}{S['capex']}-{rev30}*{I('rev_growth_lt')}/(1+{I('rev_growth_lt')})*{I('nwc_pct')})"
        w = f"Scenarios!{mid}{S['wacc']}"
        tv = f"AVERAGE({ufcf30}*(1+{I('g')})/({w}-{I('g')}),{ebitda30}*{col}${hdr_r})"
        # interim PV: reuse the mid scenario's 2027 UFCF, interpolate to this grid's UFCF30
        u27 = f"Scenarios!{mid}{S['ufcf27']}"
        pv = f"({u27}/(1+{w})^0.7+({u27}*2/3+{ufcf30}/3)/(1+{w})^1.7+({u27}/3+{ufcf30}*2/3)/(1+{w})^2.7+{ufcf30}/(1+{w})^3.7)"
        f = f"=({pv}+{tv}/(1+{w})^3.7-DCF!{D['netdebt']}-DCF!{D['debtlike']})/{I('shares')}"
        cell = v.cell(row=r, column=2 + j, value=f); cell.number_format = USD2
    r += 1
v.cell(row=r, column=1, value=f"Market: ${INV['price']:.0f}. Rows are NA margin; EMEA held at the mid-cycle scenario's margin.").font = GREY
r += 3
v.cell(row=r, column=1, value="2027E enterprise EBITDA vs price that sticks (mid-cycle levers otherwise)").font = BOLD; r += 1
v.cell(row=r, column=1, value=f"Share of ${INV['price_index_ytd']}/ton that sticks").font = BOLD
v.cell(row=r, column=2, value="$/ton").font = BOLD
v.cell(row=r, column=3, value="2027E EBITDA").font = BOLD
v.cell(row=r, column=4, value="vs 2026 guide mid").font = BOLD
v.cell(row=r, column=5, value=f"EV/EBITDA at ${INV['price']:.0f}").font = BOLD
r += 1
for pct in [0.0, 0.25, 0.5, 0.75, 1.0]:
    v.cell(row=r, column=1, value=pct).font = BLUE; v.cell(row=r, column=1).number_format = PCT
    v.cell(row=r, column=2, value=f"=A{r}*{I('price_index_ytd')}").number_format = USD
    v.cell(row=r, column=3, value=f"=Scenarios!C{S['ebitda27']}-(Scenarios!C{S['stick']}-A{r})*{I('price_index_ytd')}*{I('price_sens')}").number_format = USD
    v.cell(row=r, column=4, value=f"=C{r}/(({I('ebitda_2026_guide_lo')}+{I('ebitda_2026_guide_hi')})/2)-1").number_format = PCT
    v.cell(row=r, column=5, value=f"=(DCF!{D['mcap']}+DCF!{D['netdebt']})/C{r}").number_format = MULT
    r += 1
v.cell(row=r, column=1, value=f"Each $10/ton ≈ ${INV['price_sens'] * 10:.0f}M annualized EBITDA (CFO). This is the single assumption the stock argues about.").font = GREY

# =====================================================================================
# COMPS + SOTP
# =====================================================================================
cps = wb.create_sheet("Comps_SOTP")
setw(cps, [44, 14, 14, 14, 60])
cps["A1"] = "Trading comps and sum-of-the-parts (the separation is live: spin planned, sale process opened Oct 15, 2026)"; cps["A1"].font = H1
r = 3
header(cps, r, "Trading comps (LTM, Multiples.vc)", 5); r += 1
cps.cell(row=r, column=1, value="Company").font = BOLD; cps.cell(row=r, column=2, value="EV/EBITDA").font = BOLD; cps.cell(row=r, column=3, value="EBITDA margin").font = BOLD; r += 1
for nm, mu, mg in [("Packaging Corp of America (PKG)", I('pkg_ev_ebitda'), I('pkg_margin')), ("Smurfit Westrock (SW)", I('sw_ev_ebitda'), I('sw_margin'))]:
    cps.cell(row=r, column=1, value=nm); cps.cell(row=r, column=2, value=f"={mu}").number_format = MULT; cps.cell(row=r, column=3, value=f"={mg}").number_format = PCT
    cps.cell(row=r, column=2).font = GREEN; cps.cell(row=r, column=3).font = GREEN; r += 1
cps.cell(row=r, column=1, value=f"International Paper (IP) — 2026E, at ${INV['price']:.0f}"); cps.cell(row=r, column=2, value=f"=DCF!{D['mkt_mult26']}").number_format = MULT; cps.cell(row=r, column=3, value=f"=Model!C{R['ebitda_m']}").number_format = PCT; r += 1
cps.cell(row=r, column=1, value="IP NA segment margin 2026E"); cps.cell(row=r, column=3, value=f"=Model!C{R['ebitda_na_m']}").number_format = PCT; r += 2

header(cps, r, "Sum of the parts — 2028E mid-cycle EBITDA × segment multiple (not the 2027 price-spike year)", 5); r += 1
cps.cell(row=r, column=1, value="Segment").font = BOLD; cps.cell(row=r, column=2, value="2028E EBITDA").font = BOLD; cps.cell(row=r, column=3, value="Multiple").font = BOLD; cps.cell(row=r, column=4, value="Value").font = BOLD; cps.cell(row=r, column=5, value="Why").font = BOLD; r += 1
C = {}
cps.cell(row=r, column=1, value="Packaging Solutions NA"); cps.cell(row=r, column=2, value=f"=Model!E{R['ebitda_na']}").number_format = USD
cps.cell(row=r, column=3, value=9.0).font = BLUE; cps.cell(row=r, column=3).fill = YELLOW; cps.cell(row=r, column=3).number_format = MULT
cps.cell(row=r, column=4, value=f"=B{r}*C{r}").number_format = USD
cps.cell(row=r, column=5, value=f"Between SW ({INV['sw_ev_ebitda']}x) and PKG ({INV['pkg_ev_ebitda']}x): a pure-NA box business with 80/20 margins not yet proven.").font = GREY; C["na"] = r; r += 1
cps.cell(row=r, column=1, value="Packaging Solutions EMEA (DS Smith)"); cps.cell(row=r, column=2, value=f"=Model!E{R['ebitda_emea']}").number_format = USD
cps.cell(row=r, column=3, value=6.5).font = BLUE; cps.cell(row=r, column=3).fill = YELLOW; cps.cell(row=r, column=3).number_format = MULT
cps.cell(row=r, column=4, value=f"=B{r}*C{r}").number_format = USD
cps.cell(row=r, column=5, value="European packaging peers trade 6–7x; a sale to a strategic could clear higher — that's the Oct 15 news.").font = GREY; C["emea"] = r; r += 1
cps.cell(row=r, column=1, value="Corporate (capitalized)"); cps.cell(row=r, column=2, value=f"=Model!E{R['corp']}").number_format = USD
cps.cell(row=r, column=3, value=f"=C{C['na']}").number_format = MULT
cps.cell(row=r, column=4, value=f"=B{r}*C{r}").number_format = USD; C["corp"] = r; r += 1
cps.cell(row=r, column=1, value="Enterprise value (SOTP)").font = BOLD; cps.cell(row=r, column=4, value=f"=SUM(D{C['na']}:D{C['corp']})").number_format = USD; cps.cell(row=r, column=4).font = BOLD; C["ev"] = r; r += 1
cps.cell(row=r, column=1, value="Less: net debt + debt-like"); cps.cell(row=r, column=4, value=f"=-DCF!{D['netdebt']}-DCF!{D['debtlike']}").number_format = USD; C["nd"] = r; r += 1
cps.cell(row=r, column=1, value="Less: separation costs and dis-synergies (one-time)"); cps.cell(row=r, column=4, value=f"=-{I('sep_costs_2026')}*1.5").number_format = USD; cps.cell(row=r, column=5, value="1.5 years of the Inputs run-rate; two public companies cost more to run than one.").font = GREY; C["sep"] = r; r += 1
cps.cell(row=r, column=1, value="Equity value").font = BOLD; cps.cell(row=r, column=4, value=f"=D{C['ev']}+D{C['nd']}+D{C['sep']}").number_format = USD; cps.cell(row=r, column=4).font = BOLD; C["eq"] = r; r += 1
cps.cell(row=r, column=1, value="Per share ($)").font = BOLD; cps.cell(row=r, column=4, value=f"=D{C['eq']}/{I('shares')}").number_format = USD2; cps.cell(row=r, column=4).font = BOLD; C["ps"] = r; r += 1
cps.cell(row=r, column=1, value=f"vs. ${INV['price']:.0f} market"); cps.cell(row=r, column=4, value=f"=D{C['ps']}/{I('price')}-1").number_format = PCT; r += 2
cps.cell(row=r, column=1, value="Breakeven: EMEA multiple at which SOTP = market price").font = BLACK
cps.cell(row=r, column=4, value=f"=((DCF!{D['mcap']}-D{C['sep']}-D{C['nd']})-D{C['na']}-D{C['corp']})/B{C['emea']}").number_format = MULT
cps.cell(row=r, column=5, value="If a buyer pays more than this for DS Smith, the market is underpricing the sum.").font = GREY

# =====================================================================================
# README
# =====================================================================================
rd = wb.create_sheet("README", 0)
setw(rd, [110])
lines = [
    ("A corrugator engineer values International Paper — model notes", H1),
    ("", None),
    ("What this is: a Wall Street Prep-style DCF (unlevered FCF, WACC, terminal value two ways, comps, sum-of-the-parts) built around Damodaran's rule — write the story first, then make every number consistent with it.", BLACK),
    ("What it is not: a price target, investment advice, or a forecast. It's three stories with numbers attached, so the argument can be had line by line.", BLACK),
    ("", None),
    ("How to use it", BOLD),
    ("1. Inputs — every assumption, tagged REPORTED / GUIDANCE / MARKET / JUDGMENT with its source. Yellow cells are floor judgments: overwrite them with your own.", BLACK),
    ("2. Model — 2025A pro forma, 2026E at guidance, 2027E built as a bridge (price that sticks, volume, cost-out, inflation), 2028–30E at mid-cycle margins. Sanity rows check the dividend against cash flow.", BLACK),
    ("3. DCF — WACC build, PV of 2027–30 cash flows, terminal value as the average of Gordon growth and an exit multiple, with each cross-checked against the other. Terminal value's share of EV is shown on purpose.", BLACK),
    ("4. Scenarios — Bear / Mid-cycle / Bull side by side on the same chain; probability weights are yours.", BLACK),
    ("5. Sensitivity — the two things that decide the answer: NA mid-cycle margin × exit multiple, and how much of the 2026 price increase sticks.", BLACK),
    ("6. Comps_SOTP — PKG and SW multiples, and a sum-of-the-parts for the live separation of EMEA (DS Smith), including the EMEA multiple at which the parts equal today's price.", BLACK),
    ("", None),
    ("Conventions: blue = hardcoded input · black = formula · green = link to another sheet · yellow fill = judgment to check. No results are hardcoded; change an input and everything recalculates.", GREY),
    ("Sources are cited on the Inputs sheet next to each number (FY2025 10-K, Q2 2026 earnings release, Q1 and Q4 calls via Packaging Dive/MarketBeat, Fastmarkets RISI via Packaging Dive, Multiples.vc, Benzinga). Public information only.", GREY),
    ("Known simplifications: enterprise DCF assumes EMEA is retained (the SOTP handles the separation); taxes on EBIT with no loss carryforwards; debt-like items from the FY2025 10-K at Dec 31, 2025 balances against Q2 2026 debt and cash (pensions are only remeasured once a year); 2025A revenue split and D&A are pro forma approximations.", GREY),
]
for i, (txt, f) in enumerate(lines, 1):
    c = rd.cell(row=i, column=1, value=txt)
    if f: c.font = f
    c.alignment = Alignment(wrap_text=True, vertical="top")

# =====================================================================================
# DICTIONARY (second sheet): every term in plant language, quoting the model's own numbers.
# The numbers come from numbers.json, so run_all.py builds twice: once to compute, once to quote.
# =====================================================================================
_here = os.path.dirname(os.path.abspath(__file__))
if os.path.exists(os.path.join(_here, "numbers.json")):
    import json
    from dictionary import TITLE, INTRO, rows as dictionary_rows
    dw = wb.create_sheet("Dictionary", 1)
    setw(dw, [26, 44, 44, 40])
    dw["A1"] = TITLE; dw["A1"].font = H1
    dw["A2"] = INTRO; dw["A2"].font = GREY
    rr = 4
    for kind, cells in dictionary_rows(json.load(open(os.path.join(_here, "numbers.json")))["fmt"]):
        if kind == "section":
            header(dw, rr, cells[0], 4)
        else:
            for j, t in enumerate(cells, 1):
                c = dw.cell(row=rr, column=j, value=t)
                c.font = BOLD if (kind == "head" or j == 1) else BLACK
                c.alignment = Alignment(wrap_text=kind == "term", vertical="top")
        rr += 1

for sh in wb.worksheets:
    sh.sheet_view.showGridLines = False
    sh.freeze_panes = "B5" if sh.title in ("Model",) else None

out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "IP_valuation_model.xlsx")
wb.save(out)
# every input with its label, format, provenance and source, for post_numbers.py and the post's Inputs table
import json
json.dump({k: dict(label=l, value=v, fmt=f, prov=p, source=src, judgment=j) for k, (l, v, f, p, src, j) in LAB.items()},
          open(os.path.join(os.path.dirname(os.path.abspath(__file__)), "inputs.json"), "w"), indent=1, ensure_ascii=False)
print("saved", out)
print("Inputs rows:", {k: v for k, v in list(INP.items())[:3]}, "...")
print("Model rows:", R)
print("DCF cells:", D)
print("Scenario rows:", S)
