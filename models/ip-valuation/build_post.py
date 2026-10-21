"""Render posts/valuing-international-paper.html from its template and the model's outputs.

The template (post.tmpl, next to this script) holds the prose. Every number in it is
a placeholder filled from numbers.json, every table is drawn from the recalculated workbook or
bet.json, and every code block is the file as it ran. The build fails on an unknown placeholder,
on a leftover one, or on a chart whose data no longer fits its axis, so the text can't drift
from the model.

  {{key}}            numbers.json fmt[key]
  {{json:NAME}}      chart data (MC, DRV, BET, calc) as JSON
  {{table:NAME}}     inputs, model, dcf, comps, scenarios, grid, instruments
  {{include:FILE}}   a file from this folder, HTML-escaped (for <pre> blocks)
Run: python3 build_post.py
"""
import html, json, os, re
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(HERE))
SRC = os.path.join(HERE, "post.tmpl")  # not .html: the site is served as-is, and the template shouldn't be a page
OUT = os.path.join(ROOT, "posts", "valuing-international-paper.html")

N = json.load(open(os.path.join(HERE, "numbers.json")))
F, DATA = N["fmt"], N["data"]
bet = json.load(open(os.path.join(HERE, "bet.json")))
wb = load_workbook(os.path.join(HERE, "IP_valuation_model.xlsx"), data_only=True)
price = DATA["calc"]["price"]
SECTION = 'style="text-align:left;padding-top:12px"'


def cell_text(v, fmt):
    """A workbook value as the post shows it, following the cell's number format."""
    if v is None:
        return ""
    if isinstance(v, str):
        return html.escape(v)
    fmt = fmt or ""
    if "%" in fmt:
        return f"{v * 100:.1f}%"
    if fmt.endswith(('"x"', "\\x")):  # LibreOffice saves 0.0"x" as 0.0\x
        return f"{v:.1f}×"
    dec = 2 if ".00" in fmt else 1 if ".0" in fmt else 0
    if "$" in fmt:
        return f"({abs(v):,.{dec}f})" if v < 0 else f"${v:,.{dec}f}"
    if "#,##0" in fmt:
        return f"({abs(v):,.{dec}f})" if v < 0 else f"{v:,.{dec}f}"
    return f"{v:.1f}" if isinstance(v, float) and v != int(v) else f"{v:,.0f}" if isinstance(v, (int, float)) and abs(v) >= 1000 else str(v)


def td(c):
    """Numbers (and blanks) right-aligned; text cells plain."""
    if isinstance(c.value, str):
        return f"<td>{html.escape(c.value)}</td>"
    return f'<td class="n">{cell_text(c.value, c.number_format)}</td>'


def sheet_rows(ws, first, ncols, header_fills=("1F3A5F", "DCE6F1")):
    """Rows of a sheet as <tr>s: filled header rows become section headings, the rest label + values."""
    out = []
    for r in range(first, ws.max_row + 1):
        a = ws.cell(r, 1)
        if a.value is None and all(ws.cell(r, c).value is None for c in range(2, ncols + 1)):
            continue
        fill = a.fill.fgColor.rgb if a.fill is not None and a.fill.fill_type == "solid" else ""
        if isinstance(fill, str) and fill.endswith(header_fills):
            out.append(f'<tr><th colspan="{ncols}" {SECTION}>{html.escape(str(a.value))}</th></tr>')
            continue
        label = html.escape(str(a.value)) if a.value is not None else ""
        out.append(f"<tr><td>{label}</td>" + "".join(td(ws.cell(r, c)) for c in range(2, ncols + 1)) + "</tr>")
    return out


def table_inputs():
    ws = wb["Inputs"]
    rows = ['<tr><th>assumption</th><th>value</th><th>provenance</th><th>source / note</th></tr>']
    for r in range(5, ws.max_row + 1):
        a, b, c, e = (ws.cell(r, k) for k in (1, 2, 3, 5))
        if a.value is None:
            continue
        if b.value is None and c.value is None:
            rows.append(f'<tr><th colspan="4" {SECTION}>{html.escape(a.value)}</th></tr>')
            continue
        val = f"<td>{html.escape(b.value)}</td>" if isinstance(b.value, str) else f'<td class="n">{cell_text(b.value, b.number_format)}</td>'
        rows.append(f"<tr><td>{html.escape(a.value)}</td>{val}<td>{html.escape(str(c.value or ''))}</td><td>{html.escape(str(e.value or ''))}</td></tr>")
    return '<table class="sens xl">' + "".join(rows) + "</table>"


def table_sheet(name, first, ncols):
    return '<table class="sens xl">' + "".join(sheet_rows(wb[name], first, ncols)) + "</table>"


def table_scenarios():
    ws = wb["Scenarios"]
    rows = ["<tr><th>the chain</th><th>bear</th><th>mid-cycle</th><th>bull</th></tr>"]
    started = False
    for r in range(1, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if a == "Lever":
            started = True
            continue
        if not started or a is None:
            continue
        if isinstance(a, str) and a.startswith("Probability weight"):
            break
        if ws.cell(r, 2).value is None:
            rows.append(f'<tr><th colspan="4" {SECTION}>derived — formulas only from here down</th></tr>')
            continue
        bold = ws.cell(r, 1).font.b
        sty = ' style="font-weight:600"' if bold else ""
        cells = "".join(f'<td class="n"{sty}>{cell_text(ws.cell(r, c).value, ws.cell(r, c).number_format)}</td>' for c in (2, 3, 4))
        rows.append(f"<tr><td{sty}>{html.escape(a)}</td>{cells}</tr>")
    return '<table class="sens">' + "".join(rows) + "</table>"


def table_grid():
    g = DATA["grid"]
    ann = {round(DATA["calc"]["margin"], 4): "this model"}
    comps = json.load(open(os.path.join(HERE, "inputs.json")))
    ann.setdefault(round(comps["sw_margin"]["value"], 4), "Smurfit Westrock")
    ann.setdefault(round(comps["pkg_margin"]["value"], 4), "Packaging Corp")
    mult = lambda m: f"{m:g}×"
    rows = ["<tr><th>NA mid-cycle margin ↓ · exit multiple →</th>" + "".join(f"<th>{mult(m)}</th>" for m in g["mults"]) + "</tr>"]
    for margin, vals in g["rows"]:
        label = f"{margin * 100:.0f}%" + (f" ({ann[round(margin, 4)]})" if round(margin, 4) in ann else "")
        cells = []
        for m, v in zip(g["mults"], vals):
            cls = ' class="mkt"' if (abs(margin - g["model_margin"]) < 1e-9 and m == g["model_mult_col"]) else ' class="lo"' if v <= price else ""
            cells.append(f"<td{cls}>${v:,.0f}</td>")
        rows.append(f"<tr><td>{label}</td>" + "".join(cells) + "</tr>")
    return '<table class="sens">\n  ' + "\n  ".join(rows) + "\n</table>"


def table_instruments():
    money = lambda v: f"({abs(v):,.0f})" if v < 0 else f"${v:,.0f}"
    head = "<tr><th>instrument (per 100 shares)</th><th>capital</th><th>E[P&amp;L]</th><th>P(profit)</th><th>10th</th><th>90th</th><th>worst</th><th>on capital</th></tr>"
    rows = [f"<tr><td>{html.escape(r['name'])}</td><td class='n'>{money(r['capital'])}</td><td class='n'>{money(r['mean'])}</td><td class='n'>{r['p_profit']:.0%}</td>"
            f"<td class='n'>{money(r['p10'])}</td><td class='n'>{money(r['p90'])}</td><td class='n'>{money(r['worst'])}</td><td class='n'>{r['roc']:+.0%}</td></tr>"
            for r in bet["rows"]]
    return '<table class="sens">' + head + "".join(rows) + "</table>"


TABLES = {"inputs": table_inputs, "model": lambda: table_sheet("Model", 4, 7), "dcf": lambda: table_sheet("DCF", 4, 5),
          "comps": lambda: table_sheet("Comps_SOTP", 3, 5), "scenarios": table_scenarios, "grid": table_grid, "instruments": table_instruments}
JSONS = {"MC": lambda: json.dumps(DATA["MC"], separators=(",", ":")), "DRV": lambda: json.dumps(DATA["MC"]["drivers"]),
         "BET": lambda: json.dumps({"grid": bet["grid"], "curves": bet["curves"], "priceT": bet["priceT"], "price": bet["price"],
                                    "premiums": bet["premiums"], "c": bet["convergence"]}, separators=(",", ":")),
         "calc": lambda: json.dumps(DATA["calc"], separators=(", ", ": "))}


def render(src):
    def sub(m):
        key = m.group(1).strip()
        if key.startswith("table:"):
            return TABLES[key[6:]]()
        if key.startswith("json:"):
            return JSONS[key[5:]]()
        if key.startswith("include:"):
            return html.escape(open(os.path.join(HERE, key[8:])).read().rstrip("\n"))
        if key not in F:
            raise KeyError(f"unknown placeholder {{{{{key}}}}} — add it to post_numbers.py")
        return html.escape(F[key], quote=False)
    out = re.sub(r"\{\{(.*?)\}\}", sub, src)
    left = re.findall(r"\{\{\s*[\w:.\-]+\s*\}\}", out)
    assert not left, f"placeholders survived the build: {left[:5]}"
    return out


# the bridge chart's axis is fixed at $2,000–3,300M; fail loudly if the model outgrows it
b = DATA["bridge"]
running = [b[0], b[0] + b[1], b[0] + b[1] + b[2], b[0] + b[1] + b[2] + b[3], b[5], b[6]]
assert 2000 <= min(running) and max(running) <= 3300, f"bridge {b} no longer fits the chart's $2,000–3,300M axis"
assert b[0] + sum(b[1:5]) == b[5], f"bridge steps {b} don't add to the 2027 total"

open(OUT, "w").write(render(open(SRC).read()))
print(f"wrote {os.path.relpath(OUT, ROOT)}")
